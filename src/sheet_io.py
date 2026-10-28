"""Spreadsheet layout: A1 = N; data in columns B.. rows 1..N (row 1 is data)."""
import math
import re
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from errors import UserError


def parse_column(text: str) -> int:
    """'B' or '2' -> 2 (1-based). Column A (1) is reserved for N and rejected."""
    s = str(text).strip()
    if re.fullmatch(r"\d+", s):
        col = int(s)
    elif re.fullmatch(r"[A-Za-z]{1,3}", s):
        col = 0
        for ch in s.upper():
            col = col * 26 + (ord(ch) - ord("A") + 1)
    else:
        raise UserError(f"'{text}' is not a valid column. Use a letter (B) or a number (2).")
    if col < 2:
        raise UserError("Column A is reserved for N. Use column B (2) or later.")
    return col


def _check_file(path) -> Path:
    if not path:
        raise UserError("Choose a spreadsheet file first.")
    p = Path(path)
    if not p.is_file():
        raise UserError(f"File not found: {p}")
    if p.suffix.lower() != ".xlsx":
        raise UserError("Only .xlsx files are supported.")
    return p


def _open(path, data_only: bool):
    try:
        return load_workbook(path, data_only=data_only)
    except Exception as e:
        raise UserError(f"Could not open the spreadsheet: {e}")


def read_n(path) -> int:
    wb = _open(_check_file(path), data_only=True)
    v = wb.active["A1"].value
    wb.close()
    if v is None:
        raise UserError("Cell A1 is empty. It must hold the vector size N.")
    if isinstance(v, bool) or not isinstance(v, (int, float)) or v != int(v) or v < 1:
        raise UserError(f"Cell A1 must be a positive integer (found {v!r}).")
    return int(v)


def read_column(path, col: int, n: int) -> np.ndarray:
    wb = _open(_check_file(path), data_only=True)
    ws = wb.active
    out = np.empty(n)
    try:
        for r in range(1, n + 1):
            v = ws.cell(row=r, column=col).value
            ref = f"{get_column_letter(col)}{r}"
            if v is None:
                raise UserError(
                    f"{ref} is empty. If this cell holds a formula, open the file in Excel "
                    "and save it so the values are stored, then retry."
                )
            if isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v):
                raise UserError(f"{ref} is not a number (found {v!r}).")
            out[r - 1] = v
    finally:
        wb.close()
    return out


def column_has_data(path, col: int, n: int) -> bool:
    """True if any of rows 1..n in the column is non-empty (formulas count)."""
    wb = _open(_check_file(path), data_only=False)
    ws = wb.active
    found = any(ws.cell(row=r, column=col).value is not None for r in range(1, n + 1))
    wb.close()
    return found


def write_column(path, col: int, values) -> None:
    """Write rows 1..len(values) of one column; every other cell/formula is kept."""
    p = _check_file(path)
    wb = _open(p, data_only=False)  # not data_only: keeps formulas elsewhere
    ws = wb.active
    for r, v in enumerate(values, start=1):
        ws.cell(row=r, column=col, value=float(v))
    try:
        wb.save(p)
    except PermissionError:
        raise UserError("Could not save the file. Close it in Excel and retry.")
    finally:
        wb.close()
