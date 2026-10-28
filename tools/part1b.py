"""Build the Part I-b test workbook (N=128) and run every experiment through the app's Controller.

Usage:  python tools/part1b.py [output_dir]      (default: ./part1b_output)

Writes <output_dir>/part1b_data.xlsx (inputs in B..K, forward results in L..U, inverse
round trips in V..AE), PNG graphs for every vector, and prints a summary table.
"""
import sys
from pathlib import Path

import numpy as np
from matplotlib.figure import Figure
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

import plots  # noqa: E402
import sheet_io  # noqa: E402
import transform as tf  # noqa: E402
from controller import Controller  # noqa: E402

N = 128
n = np.arange(N)
rng = np.random.default_rng(12345)


def pulse(width):
    v = np.zeros(N)
    v[:width] = 1.0
    return v


# (name, vector) in sheet order: B, C, D, ...
VECTORS = [
    ("cos(2pi*6n/128)", np.cos(2 * np.pi * 6 * n / N)),
    ("sin(2pi*6n/128)", np.sin(2 * np.pi * 6 * n / N)),
    ("cos(2pi*9n/128)", np.cos(2 * np.pi * 9 * n / N)),
    ("sin(2pi*9n/128)", np.sin(2 * np.pi * 9 * n / N)),
    ("cos9+sin9", np.cos(2 * np.pi * 9 * n / N) + np.sin(2 * np.pi * 9 * n / N)),
    ("cos(2pi*6n/128+45deg)", np.cos(2 * np.pi * 6 * n / N + np.pi / 4)),
    ("random uniform[-1,1]", rng.uniform(-1, 1, N)),
    ("pulse width 10", pulse(10)),
    ("pulse width 5", pulse(5)),
    ("pulse width 30", pulse(30)),
]
K = len(VECTORS)
IN0, FWD0, INV0 = 2, 2 + K, 2 + 2 * K  # first column of each block


def build_workbook(path: Path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = N
    for j, (_, vec) in enumerate(VECTORS):
        for r in range(N):
            ws.cell(row=r + 1, column=IN0 + j, value=float(vec[r]))
    wb.save(path)


def save_png(fig, path: Path):
    fig.savefig(path, dpi=110)


def main():
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "part1b_output")
    out.mkdir(parents=True, exist_ok=True)
    xlsx = out / "part1b_data.xlsx"
    build_workbook(xlsx)

    figs = []
    ctl = Controller(log=lambda m: None, show_figure=figs.append)
    ctl.build(str(N))
    ctl.set_sheet(xlsx)

    rows = []
    for j, (name, _) in enumerate(VECTORS):
        cin, cfwd, cinv = (get_column_letter(c + j) for c in (IN0, FWD0, INV0))
        figs.clear()
        ctl.go("forward", cin, cfwd)
        ctl.go("inverse", cfwd, cinv)
        x = sheet_io.read_column(xlsx, IN0 + j, N)
        y = sheet_io.read_column(xlsx, FWD0 + j, N)
        xr = sheet_io.read_column(xlsx, INV0 + j, N)
        top = np.argsort(-np.abs(y))[:3]
        rows.append((name, tf.energy(x), tf.energy(y), float(np.max(np.abs(xr - x))),
                     [(int(i), float(y[i])) for i in top]))
        tag = f"{j + 1:02d}"
        save_png(plots.plot_vector(x, f"Input: {name}  (energy {tf.energy(x):.6g})"),
                 out / f"{tag}_input.png")
        save_png(plots.plot_vector(y, f"Forward transform of {name}  (energy {tf.energy(y):.6g})"),
                 out / f"{tag}_transform.png")

    # Part vii: cos(w n + 45deg) = cos45*cos(w n) - sin45*sin(w n); transform is linear
    y_cos6 = sheet_io.read_column(xlsx, FWD0 + 0, N)
    y_sin6 = sheet_io.read_column(xlsx, FWD0 + 1, N)
    y_shift = sheet_io.read_column(xlsx, FWD0 + 5, N)
    ident_err = float(np.max(np.abs(y_shift - np.sqrt(0.5) * (y_cos6 - y_sin6))))
    # Part vi: transform of the sum equals the sum of the transforms
    y_c9 = sheet_io.read_column(xlsx, FWD0 + 2, N)
    y_s9 = sheet_io.read_column(xlsx, FWD0 + 3, N)
    y_sum = sheet_io.read_column(xlsx, FWD0 + 4, N)
    lin_err = float(np.max(np.abs(y_sum - (y_c9 + y_s9))))

    print(f"N = {N}; workbook: {xlsx}\n")
    print(f"{'vector':26s} {'E(in)':>10s} {'E(out)':>10s} {'roundtrip err':>14s}  top-3 output bins (index: value)")
    for name, ex, ey, rt, top in rows:
        tops = ", ".join(f"{i}: {v:+.4f}" for i, v in top)
        print(f"{name:26s} {ex:10.5f} {ey:10.5f} {rt:14.2e}  {tops}")
    print(f"\nLinearity check (part vi):  max |T(cos9+sin9) - (T cos9 + T sin9)| = {lin_err:.2e}")
    print(f"45-degree identity (part vii): max |T(cos(wn+pi/4)) - sqrt(1/2)(T cos - T sin)| = {ident_err:.2e}")
    print(f"PNG graphs written to {out}")


if __name__ == "__main__":
    main()
