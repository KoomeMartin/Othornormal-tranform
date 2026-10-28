import numpy as np
import pytest
from openpyxl import Workbook, load_workbook

import sheet_io as sio
from errors import UserError


@pytest.mark.parametrize("text,col", [("B", 2), ("b", 2), ("2", 2), ("aa", 27), (" C ", 3)])
def test_parse_column_ok(text, col):
    assert sio.parse_column(text) == col


@pytest.mark.parametrize("text", ["A", "1", "0", "", "B2", "-3"])
def test_parse_column_bad(text):
    with pytest.raises(UserError):
        sio.parse_column(text)


def make_sheet(tmp_path, n=4):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = n
    for r in range(1, 5):
        ws.cell(row=r, column=2, value=r)
    p = tmp_path / "t.xlsx"
    wb.save(p)
    return p


def test_read_n_and_column(tmp_path):
    p = make_sheet(tmp_path)
    assert sio.read_n(p) == 4
    assert np.array_equal(sio.read_column(p, 2, 4), [1, 2, 3, 4])


def test_read_n_bad(tmp_path):
    p = make_sheet(tmp_path, n=2.5)
    with pytest.raises(UserError):
        sio.read_n(p)


def test_blank_and_text_cells_named(tmp_path):
    p = make_sheet(tmp_path)
    with pytest.raises(UserError, match="C1"):
        sio.read_column(p, 3, 4)
    wb = load_workbook(p)
    wb.active["B3"] = "abc"
    wb.save(p)
    with pytest.raises(UserError, match="B3"):
        sio.read_column(p, 2, 4)


def test_write_preserves_formulas_and_other_columns(tmp_path):
    p = make_sheet(tmp_path)
    wb = load_workbook(p)
    wb.active["D1"] = "=B1*2"
    wb.save(p)
    sio.write_column(p, 3, [9.0, 8.0, 7.0, 6.0])
    ws = load_workbook(p).active
    assert ws["D1"].value == "=B1*2"
    assert [ws.cell(row=r, column=3).value for r in range(1, 5)] == [9, 8, 7, 6]
    assert [ws.cell(row=r, column=2).value for r in range(1, 5)] == [1, 2, 3, 4]
    assert ws["A1"].value == 4


def test_column_has_data(tmp_path):
    p = make_sheet(tmp_path)
    assert sio.column_has_data(p, 2, 4)
    assert not sio.column_has_data(p, 3, 4)


def test_file_checks(tmp_path):
    with pytest.raises(UserError):
        sio.read_n("")
    with pytest.raises(UserError):
        sio.read_n(tmp_path / "missing.xlsx")
    bad = tmp_path / "x.csv"
    bad.write_text("1")
    with pytest.raises(UserError):
        sio.read_n(bad)
