import numpy as np
import pytest
from openpyxl import Workbook, load_workbook

from controller import Controller
from errors import UserError


def make(confirm=lambda m: True):
    logs, figs = [], []
    c = Controller(logs.append, figs.append, confirm=confirm)
    return c, logs, figs


def sheet(tmp_path, n=8, values=None):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = n
    vals = values if values is not None else np.cos(2 * np.pi * 1 * np.arange(n) / n)
    for r in range(n):
        ws.cell(row=r + 1, column=2, value=float(vals[r]))
    p = tmp_path / "s.xlsx"
    wb.save(p)
    return p


@pytest.mark.parametrize("bad", ["6", "0", "1", "-4", "abc", "3.5", "", "8192"])
def test_build_rejects(bad):
    c, _, _ = make()
    with pytest.raises(UserError):
        c.build(bad)
    assert c.matrix is None


@pytest.mark.parametrize("ok", ["2", "4", "128"])
def test_build_accepts(ok):
    c, logs, _ = make()
    assert c.build(ok) == int(ok)
    assert "Built" in logs[-1]


def test_failed_rebuild_keeps_old_matrix():
    c, _, _ = make()
    c.build("8")
    with pytest.raises(UserError):
        c.build("7")
    assert c.n == 8 and c.matrix.shape == (8, 8)


def test_verify_logs_pass_and_explanation():
    c, logs, _ = make()
    c.build("4")
    c.verify()
    text = "\n".join(logs)
    assert "PASS" in text and "prove the rows are orthonormal" in text


def test_verify_needs_matrix():
    c, _, _ = make()
    with pytest.raises(UserError):
        c.verify()


def test_plot_row_range():
    c, _, _ = make()
    c.build("8")
    for bad in ["8", "-1", "x", ""]:
        with pytest.raises(UserError):
            c.plot_row(bad)


def test_plot_row_logs_summary_and_shows_figure():
    c, logs, figs = make()
    c.build("128")
    c.plot_row("5")
    assert len(figs) == 1
    assert "5 sign change(s)" in logs[-1] and "2.5 cycles" in logs[-1]


def test_plot_row_needs_matrix():
    c, _, _ = make()
    with pytest.raises(UserError):
        c.plot_row("0")


def test_go_forward_inverse_round_trip(tmp_path):
    p = sheet(tmp_path)
    c, logs, figs = make()
    c.build("8")
    c.set_sheet(p)
    c.go("forward", "B", "C")
    c.go("inverse", "C", "D")
    ws = load_workbook(p).active
    b = [ws.cell(row=r, column=2).value for r in range(1, 9)]
    d = [ws.cell(row=r, column=4).value for r in range(1, 9)]
    assert np.allclose(b, d, atol=1e-12)
    assert len(figs) == 2


def test_go_energy_preserved(tmp_path):
    p = sheet(tmp_path)
    c, logs, _ = make()
    c.build("8")
    c.set_sheet(p)
    c.go("forward", "b", "c")
    assert any("ratio = 1" in l for l in logs)


def test_go_n_mismatch(tmp_path):
    p = sheet(tmp_path, n=8)
    c, _, _ = make()
    c.build("4")
    c.set_sheet(p)
    with pytest.raises(UserError, match="N=8"):
        c.go("forward", "B", "C")


def test_go_rejects_column_a_and_missing_sheet(tmp_path):
    p = sheet(tmp_path)
    c, _, _ = make()
    c.build("8")
    with pytest.raises(UserError):
        c.go("forward", "B", "C")  # no sheet chosen
    c.set_sheet(p)
    with pytest.raises(UserError):
        c.go("forward", "B", "A")
    with pytest.raises(UserError):
        c.go("forward", "A", "C")


def test_go_overwrite_declined_writes_nothing(tmp_path):
    p = sheet(tmp_path)
    c, logs, figs = make(confirm=lambda m: False)
    c.build("8")
    c.set_sheet(p)
    c.go("forward", "B", "B")
    assert figs == [] and any("cancelled" in l for l in logs)
    ws = load_workbook(p).active
    assert ws["B1"].value == pytest.approx(1.0)


def test_go_blank_cell_named(tmp_path):
    p = sheet(tmp_path)
    c, _, _ = make()
    c.build("8")
    c.set_sheet(p)
    with pytest.raises(UserError, match="C1"):
        c.go("forward", "C", "D")


def test_graph_column_energy(tmp_path):
    p = sheet(tmp_path, values=np.ones(8))
    c, logs, figs = make()
    c.build("8")
    c.set_sheet(p)
    c.graph_column("B")
    assert any("energy = 8" in l for l in logs) and len(figs) == 1
